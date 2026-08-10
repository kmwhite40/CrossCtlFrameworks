"""Spreadsheet parser.

Each worksheet becomes a page and the sheet name becomes both the section and the
table id — a control matrix and an asset inventory in the same workbook are
different evidence, and flattening them together would lose that. Row 1 is
treated as the header row, so every value below it carries the column label that
makes it legible.

Blank cells are skipped but never shift coordinates: ``col_index`` is the true
spreadsheet column, so a citation points at the cell a reader would actually
find.
"""

from __future__ import annotations

import io
import logging

import openpyxl

from .base import ParsedBlock, ParsedCell, ParsedDocument, ParsedPage

log = logging.getLogger(__name__)

PARSER_NAME = "xlsx"
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_xlsx(data: bytes, filename: str, media_type: str = MEDIA_TYPE) -> ParsedDocument:
    """Parse a .xlsx into one page per worksheet of header-labelled cell blocks."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # a malformed workbook must not kill the run
        log.warning("prep.parse.xlsx_failed", extra={"filename": filename, "error": str(exc)})
        return ParsedDocument(
            filename=filename, media_type=media_type, parser_name=PARSER_NAME,
            error=f"could not open workbook: {exc}",
        )

    pages: list[ParsedPage] = []
    try:
        for page_number, sheet in enumerate(workbook.worksheets, start=1):
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                pages.append(ParsedPage(page_number=page_number, section_title=sheet.title))
                continue
            headers = [_as_text(v) for v in rows[0]]
            blocks: list[ParsedBlock] = []
            for row_index, row in enumerate(rows):
                cells = [
                    ParsedCell(
                        text=_as_text(value),
                        row_index=row_index,
                        col_index=col_index,
                        header=headers[col_index] if col_index < len(headers) else None,
                    )
                    for col_index, value in enumerate(row)
                ]
                for cell in cells:
                    if not cell.text:
                        continue
                    blocks.append(
                        ParsedBlock(
                            block_id=f"{sheet.title}r{cell.row_index}c{cell.col_index}",
                            block_type="sheet_cell",
                            text=cell.text,
                            heading_path=[sheet.title],
                            row_index=cell.row_index,
                            col_index=cell.col_index,
                            table_id=sheet.title,
                            # The header row labels itself otherwise, which is noise.
                            cell_label=cell.header if row_index > 0 else None,
                            cells=cells,
                        )
                    )
            pages.append(
                ParsedPage(page_number=page_number, blocks=blocks, section_title=sheet.title)
            )
    finally:
        workbook.close()

    return ParsedDocument(
        filename=filename, media_type=media_type, parser_name=PARSER_NAME, pages=pages
    )
