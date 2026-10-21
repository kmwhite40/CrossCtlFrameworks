"""A policy version as a Word document.

The body is the version's own Markdown (headings, paragraphs, bullets), so the
document says exactly what the hashed text says. The control block records the
version, its status and its SHA-256, so a printed copy can be checked against
the record. A draft is marked DRAFT on every page; an unset organization-defined
value stays visible as "[ORGANIZATION-DEFINED: key]".
"""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt, RGBColor

from ..models import Policy, PolicyVersion

_DRAFT_RED = RGBColor(0xB9, 0x1C, 0x1C)


def to_docx(policy: Policy, version: PolicyVersion, *, organization: str | None) -> bytes:
    doc = Document()
    in_force = policy.active_version_id == version.id and policy.status == "active"
    status = "IN FORCE" if in_force else ("RETIRED" if policy.status == "retired" else "DRAFT")

    header = doc.sections[0].header.paragraphs[0]
    header.text = f"{organization or ''}  |  {policy.name}  |  v{version.version}  |  {status}"
    if not in_force:
        for run in header.runs:
            run.font.color.rgb = _DRAFT_RED
            run.font.bold = True

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in (
        ("Document", policy.name),
        ("Version", version.version),
        ("Status", status),
        ("Effective", str(version.effective_on) if version.effective_on else "—"),
        ("Review", policy.review_frequency or "—"),
        ("Next review", str(policy.next_review_on) if policy.next_review_on else "—"),
        ("SHA-256", version.body_sha256 or "—"),
    ):
        row = table.add_row().cells
        row[0].text, row[1].text = label, value
    doc.add_paragraph()

    if version.body is None:
        doc.add_paragraph(f"The document is held externally: {version.uri or '—'}")
    else:
        for block in version.body.split("\n\n"):
            text = block.strip()
            if not text:
                continue
            if text.startswith("# "):
                doc.add_heading(text[2:].strip(), level=0)
            elif text.startswith("## "):
                doc.add_heading(text[3:].strip(), level=1)
            elif all(line.startswith("- ") for line in text.splitlines()):
                for line in text.splitlines():
                    doc.add_paragraph(line[2:].strip(), style="List Bullet")
            else:
                p = doc.add_paragraph(" ".join(line.strip() for line in text.splitlines()))
                for run in p.runs:
                    run.font.size = Pt(11)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
