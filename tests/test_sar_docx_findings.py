"""Every finding a SAR objective can hold renders under a visible name.

``ccf.assessment.sar`` renders ``AssessmentControlResult.objective_findings``
into the .docx an assessor is handed. That column has two producers with two
vocabularies -- the seeder and the assessor form write
``ccf.assessment.seed.FINDINGS``, the assessment engine writes
``ccf.models_assessment_engine.OBJECTIVE_VERDICTS`` on acceptance -- and since
the form offers the union of both, an assessor can select either.

The renderer's label map held only the first vocabulary and defaulted to the
empty string, so ``not_satisfied`` and ``insufficient_evidence`` printed as::

    [] the policy is reviewed at a defined frequency;

A failure, rendered nameless, in a federal deliverable. These tests pin the
two properties that stop it recurring: every member of either vocabulary has a
non-empty label, and NO value can render as nothing.

Pure rendering -- no DB.
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from ccf.assessment.sar import OBJECTIVE_FINDINGS, generate_sar_docx
from ccf.assessment.seed import FINDINGS
from ccf.models_assessment_engine import OBJECTIVE_VERDICTS

_VOCABULARY: tuple[str, ...] = tuple(sorted({*FINDINGS, *OBJECTIVE_VERDICTS}))

_META = {
    "customer_name": "Acme Corp",
    "system_name": "Acme System",
    "assessment_name": "sar-label-assessment",
    "kind": "self",
    "assessor": "assessor@example.com",
    "period": "2026-01-01 to ongoing",
    "date": "08/11/2026",
}
_SUMMARY = {"total": 1, "by_finding": dict.fromkeys(FINDINGS, 0), "score": 110}

_OBJECTIVE_TEXT = "the policy is reviewed at a defined frequency;"


def _render(part_finding: str, *, control_finding: str = "satisfied") -> str:
    """The Part row's rendered text for one objective holding ``part_finding``."""
    data = generate_sar_docx(
        _META,
        _SUMMARY,
        [
            {
                "domain": "AC",
                "control_id": "AC.L2-3.1.1",
                "nist_id": "3.1.1",
                "title": "Authorized Access Control",
                "finding": control_finding,
                "objective_findings": [
                    {"label": "a", "text": _OBJECTIVE_TEXT, "finding": part_finding}
                ],
            }
        ],
    )
    doc = Document(io.BytesIO(data))
    for table in doc.tables:
        for row in table.rows:
            if row.cells[0].text.startswith("Part ["):
                return row.cells[1].text
    raise AssertionError("no Part row rendered")


def _bracketed(cell_text: str) -> str:
    """The label the renderer put in the leading ``[...]``."""
    assert cell_text.startswith("["), cell_text
    return cell_text[1 : cell_text.index("]")]


# --- every value the column can hold gets a name -----------------------------


def test_the_renderer_covers_both_producers_vocabularies() -> None:
    """Derived from the constants, so a future member needs no second edit."""
    assert set(OBJECTIVE_FINDINGS) == set(_VOCABULARY)


@pytest.mark.parametrize("finding", _VOCABULARY)
def test_every_vocabulary_member_renders_a_visible_label(finding: str) -> None:
    text = _render(finding)
    label = _bracketed(text)
    assert label.strip(), f"{finding!r} rendered a nameless determination: {text!r}"
    assert _OBJECTIVE_TEXT in text


@pytest.mark.parametrize("finding", ("not_satisfied", "insufficient_evidence"))
def test_the_two_engine_only_verdicts_render_their_own_names(finding: str) -> None:
    """The regression proper. Both printed ``[]`` before the fix, and neither
    may be renamed onto a determination it is not."""
    label = _bracketed(_render(finding))
    assert label == finding.replace("_", " ").title()


# --- an unmapped value is loud, not blank, and asserts nothing ---------------


@pytest.mark.parametrize(
    "finding", ("legacy_import_value", "partially_satisfied", "SATISFIED", "  ")
)
def test_an_unmapped_finding_renders_visibly(finding: str) -> None:
    label = _bracketed(_render(finding))
    assert label.strip(), f"{finding!r} rendered as nothing"


@pytest.mark.parametrize("finding", ("legacy_import_value", "partially_satisfied"))
def test_an_unmapped_finding_does_not_assert_a_determination(finding: str) -> None:
    """It must not be silently read as any known determination, and it must
    carry the raw value so the document stays traceable to the row."""
    label = _bracketed(_render(finding))
    assert finding in label, f"the raw value was dropped: {label!r}"
    assert label not in set(OBJECTIVE_FINDINGS)
    for known in OBJECTIVE_FINDINGS:
        assert label != known.replace("_", " ").title()


def test_a_missing_or_blank_finding_is_not_dressed_as_a_determination() -> None:
    label = _bracketed(_render(""))
    assert label.strip()
    for known in OBJECTIVE_FINDINGS:
        assert label != known.replace("_", " ").title()


# --- existing output is unchanged where it was already a label ---------------


@pytest.mark.parametrize(
    ("finding", "expected"),
    (
        ("satisfied", "Satisfied"),
        ("other_than_satisfied", "Other Than Satisfied"),
        ("not_applicable", "Not Applicable"),
        ("not_assessed", "Not Assessed"),
    ),
)
def test_the_four_pre_existing_labels_are_byte_identical(finding: str, expected: str) -> None:
    """Deriving the map from the constants must not reword a label that was
    already correct -- the only intended movement in rendered output is a
    blank becoming a name."""
    assert _render(finding) == f"[{expected}] {_OBJECTIVE_TEXT}"


@pytest.mark.parametrize("finding", ("satisfied", "other_than_satisfied", "not_applicable"))
def test_the_control_grain_finding_row_is_unchanged(finding: str) -> None:
    data = generate_sar_docx(
        _META,
        _SUMMARY,
        [
            {
                "domain": "AC",
                "control_id": "AC.L2-3.1.1",
                "nist_id": "3.1.1",
                "title": "Authorized Access Control",
                "finding": finding,
                "objective_findings": [],
            }
        ],
    )
    doc = Document(io.BytesIO(data))
    rows = [
        r.cells[1].text
        for t in doc.tables
        for r in t.rows
        if r.cells[0].text == "Finding"
    ]
    assert rows == [finding.replace("_", " ").title()]
